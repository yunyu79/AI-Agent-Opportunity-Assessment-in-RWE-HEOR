"""HTTP smoke tests: real PDF upload/parser, demo, preview, edited exports."""
import base64
import io
import json
import re
import threading
import unittest
import urllib.error
import urllib.request
from unittest.mock import patch

from openpyxl import load_workbook
from http.server import ThreadingHTTPServer
import app
from core import demo_pdf, demo_result


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f'http://127.0.0.1:{cls.server.server_port}'

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def get(self, path):
        return urllib.request.urlopen(self.url + path)

    def post(self, path, data, token=app.TOKEN):
        req = urllib.request.Request(self.url + path, data=json.dumps(data).encode(),
             headers={'Content-Type': 'application/json', 'X-App-Token': token})
        return urllib.request.urlopen(req)

    def test_demo_page_preview_edited_excel_and_audit(self):
        with self.get('/') as res:
            self.assertIn(app.TOKEN, res.read().decode())
        with self.post('/api/demo', {}) as res: run = json.load(res)
        paper = run['papers'][0]
        with self.get(f'/api/page?run={run["run"]}&paper={paper["paper_id"]}&page=2') as res:
            self.assertTrue(res.read().startswith(b'\x89PNG'))
        rows = paper['rows']
        rows[-1]['value'] = '1.3'
        rows[-1]['review_note'] = 'Changed for export smoke test.'
        payload = {'run': run['run'], 'rows': rows}
        with self.post('/api/export', payload) as res:
            wb = load_workbook(io.BytesIO(res.read()))
            self.assertEqual(wb['Outcomes']['D4'].value, '1.3')
            self.assertEqual(wb['Audit edits'].max_row, 3)
        with self.post('/api/audit', payload) as res:
            audit = json.load(res)
            self.assertEqual(audit['papers'][0]['raw_output']['outcomes'][-1]['value'], '1.2')
            self.assertEqual(audit['reviewed_rows'][-1]['value'], '1.3')

    def test_real_upload_route_with_mocked_provider(self):
        with patch('core.extract', return_value=(demo_result(), {'model': 'mock', 'usage': {}})):
            with self.post('/api/extract', {'files': [{'name': 'upload.pdf', 'data': base64.b64encode(demo_pdf()).decode()}], 'model': 'test-model'}) as res:
                body = json.load(res)
        self.assertEqual(len(body['papers']), 1)
        self.assertFalse(body['papers'][0]['demo'])
        self.assertEqual(body['errors'], [])

    def test_invalid_upload_returns_per_file_error(self):
        with self.post('/api/extract', {'files': [{'name': 'bad.pdf', 'data': base64.b64encode(b'not a pdf').decode()}]}) as res:
            body = json.load(res)
        self.assertIsNone(body['run'])
        self.assertEqual(len(body['errors']), 1)

    def test_csrf_and_rebinding_guards(self):
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.post('/api/demo', {}, token='bad-token')
        self.assertEqual(error.exception.code, 403)
        req = urllib.request.Request(self.url + '/', headers={'Host': 'evil.example'})
        with self.assertRaises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(req)
        self.assertEqual(error.exception.code, 403)


if __name__ == '__main__': unittest.main()
