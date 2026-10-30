import copy
import io
import json
import unittest
from unittest.mock import patch
from urllib.error import HTTPError

import pymupdf
from openpyxl import load_workbook
from core import (COMMON, apply_edits, demo_pdf, demo_result, extract, make_rows,
                  paper_record, parse_pdf, validate, workbook)


class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.paper = paper_record('synthetic.pdf', demo_pdf(), demo=True)

    def test_demo_has_traceable_multi_group_outcomes(self):
        rows = self.paper['rows']
        self.assertEqual(len(rows), 10)
        self.assertTrue(all(r['validation']['quote_match'] for r in rows))
        self.assertEqual(len([r for r in rows if r['section'] == 'outcomes']), 3)
        self.assertTrue(all(r['status'] == 'Needs review' for r in rows))

    def test_invalid_or_wrong_page_quote_rejected(self):
        row = copy.deepcopy(self.paper['rows'][-1])
        row['page'] = 1
        self.assertFalse(validate(row, self.paper['pages'])['quote_match'])
        row['page'] = 300
        self.assertIn('Missing or invalid PDF page number.', validate(row, self.paper['pages'])['issues'])
        row['page'] = 2
        row['source_quote'] = 'This quotation was invented.'
        self.assertFalse(validate(row, self.paper['pages'])['quote_match'])

    def test_missing_cost_metadata_flagged(self):
        row = copy.deepcopy(self.paper['rows'][-3])
        row['currency'] = None
        self.assertIn('Cost currency is unspecified.', validate(row, self.paper['pages'])['issues'])

    def test_missing_study_field_is_not_reported_not_invented(self):
        result = demo_result()
        result['study'] = []
        rows = make_rows(result, 'demo', self.paper['pages'])
        absent = [r for r in rows if r['section'] == 'study']
        self.assertEqual(len(absent), 7)
        self.assertTrue(all(r['value'] is None and r['status'] == 'Not reported' for r in absent))

    def test_scanned_pdf_and_invalid_pdf_fail(self):
        doc = pymupdf.open()
        doc.new_page()
        with self.assertRaisesRegex(ValueError, 'text layer'):
            parse_pdf(doc.tobytes())
        doc.close()
        with self.assertRaises(ValueError): parse_pdf(b'not a pdf')

    def test_export_preserves_edits_original_and_uses_text_cells(self):
        edits = copy.deepcopy(self.paper['rows'])
        edits[0]['value'] = '=HYPERLINK("https://example.com")'
        edits[0]['page'] = 2
        rows, audit = apply_edits([self.paper], edits)
        self.assertNotEqual(rows[0]['value'], self.paper['rows'][0]['value'])
        self.assertFalse(rows[0]['validation']['quote_match'])
        self.assertEqual(len(audit), 2)
        raw = workbook([self.paper], rows, audit)
        wb = load_workbook(io.BytesIO(raw), data_only=False)
        self.assertEqual(wb.sheetnames, ['Study information', 'Outcomes', 'Audit edits', 'Run metadata'])
        self.assertEqual(wb['Study information']['D2'].data_type, 's')
        self.assertEqual(wb['Study information']['D2'].value, edits[0]['value'])
        self.assertEqual(wb['Outcomes'].max_row, 4)

    def test_export_cannot_reassign_or_omit_rows(self):
        with self.assertRaises(ValueError): apply_edits([self.paper], [])
        edits = copy.deepcopy(self.paper['rows'])
        edits[0]['id'] = 'invented'
        with self.assertRaises(ValueError): apply_edits([self.paper], edits)

    def test_mocked_live_response_uses_structured_output(self):
        class Response(io.BytesIO):
            def __enter__(self): return self
            def __exit__(self, *_): self.close()
        response = {'status': 'completed', 'output': [{'content': [{'type': 'output_text', 'text': json.dumps(demo_result())}]}], 'usage': {'input_tokens': 100}}
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-only'}), patch('urllib.request.urlopen', return_value=Response(json.dumps(response).encode())) as call:
            result, meta = extract(self.paper['pages'], 'test-model')
            request = call.call_args.args[0]
            payload = json.loads(request.data)
            self.assertTrue(payload['text']['format']['strict'])
            self.assertFalse(payload['store'])
            self.assertEqual(len(make_rows(result, 'live', self.paper['pages'])), 10)
            self.assertEqual(meta['usage']['input_tokens'], 100)

    def test_incomplete_and_refusal_are_not_accepted(self):
        for response in [{'status': 'incomplete'}, {'status': 'completed', 'output': [{'content': [{'type': 'refusal'}]}]}]:
            class Response(io.BytesIO):
                def __enter__(self): return self
                def __exit__(self, *_): pass
            with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-only'}), patch('urllib.request.urlopen', return_value=Response(json.dumps(response).encode())):
                with self.assertRaises(ValueError): extract(self.paper['pages'], 'test-model')


if __name__ == '__main__': unittest.main()
